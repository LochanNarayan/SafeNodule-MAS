"""Build the Word (.docx) version of the article by converting the generated
LaTeX file, so the two stay in lock-step. This is the primary deliverable.

    python build_paper_tex.py     # writes paper/enhanced_lungnoduleagent.tex
    python build_article.py       # writes output/Enhanced_LungNoduleAgent_Article.docx

The converter is written specifically for the constructs build_paper_tex.py
emits (numbered sections/subsections, full-grid tables, figures, boxed
algorithms marked with ALGOBOX sentinels, an adjustwidth abstract, a static
thebibliography, numeric \\cite{} citations). It is not a general TeX->docx tool.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).parent
TEX = ROOT / "paper" / "enhanced_lungnoduleagent.tex"
FIGS = ROOT / "paper" / "figures"
OUT = ROOT / "output" / "Enhanced_LungNoduleAgent_Article.docx"

BODY_SIZE = 12

# --- reference numbering / short author labels from the paper generator ---
try:
    from build_paper_tex import REFS
except Exception:
    REFS = []
REF_NUMBER = {k: i + 1 for i, (k, _lab, _txt) in enumerate(REFS)}


# ----------------------------------------------------------------------------- #
#  inline LaTeX -> plain text
# ----------------------------------------------------------------------------- #
_MATH = [
    (r"\\gets", "\u2190"), (r"\\varnothing", "\u2205"), (r"\\varepsilon", "\u03b5"),
    (r"\\geq", "\u2265"), (r"\\leq", "\u2264"), (r"\\ge\b", "\u2265"), (r"\\le\b", "\u2264"),
    (r"\\Delta", "\u0394"), (r"\\alpha", "\u03b1"), (r"\\tau", "\u03c4"),
    (r"\\theta", "\u03b8"), (r"\\ell", "\u2113"), (r"\\in\b", "\u2208"),
    (r"\\pm", "\u00b1"), (r"\\times", "\u00d7"), (r"\\approx", "\u2248"),
    (r"\\delta", "\u03b4"), (r"\\phi", "\u03c6"), (r"\\gamma", "\u03b3"),
    (r"\\wedge", "\u2227"), (r"\\vee", "\u2228"), (r"\\forall", "\u2200"),
    (r"\\mathcal\{([^}]*)\}", r"\1"),
    (r"\\cup", "\u222a"), (r"\\cap", "\u2229"), (r"\\neq", "\u2260"),
    (r"\\propto", "\u221d"), (r"\\bigg?(?=[^A-Za-z])", ""), (r"\\!", ""),
    (r"\\to\b", "\u2192"), (r"\\rightarrow", "\u2192"),
    (r"\\sum", "\u03a3"), (r"\\log", "log"), (r"\\exp", "exp"), (r"\\arg\\max", "argmax"),
    (r"\\mathcal\{R\}", "R"), (r"\\bar P", "P\u0304"), (r"\\tilde P", "P\u0303"),
    (r"\\tilde p", "p\u0303"), (r"\\bar p", "p\u0304"),
    (r"\\hat y", "\u0177"), (r"\\hphantom\{[^}]*\}", ""),
    (r"\\mathrm\{([^}]*)\}", r"\1"), (r"\\textrm\{([^}]*)\}", r"\1"),
    (r"\\text\{([^}]*)\}", r"\1"), (r"\\textstyle", ""), (r"\\textsf\{([^}]*)\}", r"\1"),
    (r"\{=\}", "="), (r"\\,", "\u2009"), (r"\\;", " "), (r"\\ ", " "),
    (r"\\big\}", "}"), (r"\\big\{", "{"), (r"\\big\/", "/"),
    (r"_\{([^}]*)\}", r"_\1"), (r"_(\w)", r"_\1"), (r"\^\{([^}]*)\}", r"\1"),
    (r"\\{", "{"), (r"\\}", "}"),
]


def _resolve_cites(s: str) -> str:
    """\\cite{k1,k2} -> "[N1,N2]" using the reference's position in REFS."""
    def rep(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        nums = sorted(REF_NUMBER.get(k, 0) for k in keys)
        return "[" + ",".join(str(n) for n in nums) + "]"
    return re.sub(r"\\cite\{([^}]+)\}", rep, s)


def _bib_clean(s: str) -> str:
    """Clean a bibliography entry but keep <i>...</i> markers for italic runs."""
    s = s.replace("\\&", "&").replace("\\%", "%")
    s = s.replace("\\ldots", "\u2026").replace("~", " ")
    s = re.sub(r"\\textit\{([^{}]*)\}", r"<i>\1</i>", s)
    s = re.sub(r"\\[A-Za-z]+\*?", "", s)
    s = s.replace("{", "").replace("}", "").replace("\\", "")
    return re.sub(r"\s+", " ", s).strip()


def detex(s: str) -> str:
    s = _resolve_cites(s)
    s = re.sub(r"\\(emph|textbf|textit|texttt|textsc)\{([^{}]*)\}", r"\2", s)
    s = re.sub(r"\\(ref|label|eqref)\{[^}]*\}", "", s)
    s = re.sub(r"\\textsuperscript\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\enquote\{([^}]*)\}",
               lambda m: "\u201c" + m.group(1) + "\u201d", s)
    s = s.replace("\\%", "%").replace("\\&", "&").replace("\\#", "#")
    s = s.replace("\\ldots", "\u2026").replace("\\dots", "\u2026")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("~", "\u00a0")
    for pat, rep in _MATH:
        s = re.sub(pat, rep, s)
    s = s.replace("$", "")
    s = re.sub(r"\\[A-Za-z]+\*?", "", s)          # drop any leftover commands
    s = s.replace("{", "").replace("}", "").replace("\\", "")
    s = re.sub(r"[ \t]+", " ", s).strip()
    return s


# ----------------------------------------------------------------------------- #
#  docx helpers
# ----------------------------------------------------------------------------- #
def font(run, size=BODY_SIZE, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def para(doc, text, *, size=BODY_SIZE, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0.2,
         before=0, after=6, italic=False, bold=False, first_indent=True,
         l_indent=0.0, r_indent=0.0, line=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if first_indent and indent:
        pf.first_line_indent = Inches(indent)
    if l_indent:
        pf.left_indent = Inches(l_indent)
    if r_indent:
        pf.right_indent = Inches(r_indent)
    font(p.add_run(text), size=size, italic=italic, bold=bold)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.line_spacing = 1.0
    if level == 1:
        pf.space_before = Pt(16); pf.space_after = Pt(8)
        font(p.add_run(text), size=16, bold=True)
    elif level == 2:
        pf.space_before = Pt(12); pf.space_after = Pt(6)
        font(p.add_run(text), size=13, bold=True, italic=True)
    else:
        pf.space_before = Pt(10); pf.space_after = Pt(4)
        font(p.add_run(text), size=12, bold=True, italic=True)
    return p


def _set_cell_border(cell, sz=5):
    """Full grid: single rule on all four sides of the cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")


def _shade_cell(cell, hex_color="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table(doc, caption, headers, rows):
    t = doc.add_table(rows=1, cols=max(len(headers), 1))
    t.autofit = True
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        font(cell.paragraphs[0].add_run(h), size=10.5, bold=True)
        _set_cell_border(cell)
        _shade_cell(cell)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            if j >= len(cells):
                break
            cells[j].text = ""
            font(cells[j].paragraphs[0].add_run(str(val)), size=10.5)
            _set_cell_border(cells[j])
    if caption:
        c = doc.add_paragraph()
        c.paragraph_format.space_before = Pt(4)
        c.paragraph_format.space_after = Pt(8)
        font(c.add_run(caption), size=11, italic=True)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_figure(doc, png: Path, caption: str, width=6.3):
    if png.exists():
        doc.add_picture(str(png), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)
    font(cap.add_run(caption), size=11, italic=True)


def add_algorithm(doc, caption, inputs, outputs, steps):
    """Framed, plainly numbered procedure box matching algo_box() in the LaTeX
    generator: a bordered one-cell table holding the Input/Output lines and the
    numbered steps, with the caption ("Algorithm N: Title") below the box."""
    t = doc.add_table(rows=1, cols=1)
    t.autofit = True
    cell = t.rows[0].cells[0]
    cell.text = ""
    _set_cell_border(cell, sz=8)
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "120"); node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)

    first = True
    for label, text in (("Input", inputs), ("Output", outputs)):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(6)
        font(p.add_run(f"{label}: "), size=11, bold=True)
        font(p.add_run(text), size=11)

    for i, item in enumerate(steps, 1):
        level, s = item if isinstance(item, tuple) else (0, item)
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.45 + 0.20 * level)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        font(p.add_run(f"{i}."), size=10.5)
        font(p.add_run("\u2002" * (1 + 2 * level)), size=10.5)
        for seg in re.split(r"(<b>.*?</b>)", s):
            if not seg:
                continue
            bold = seg.startswith("<b>")
            font(p.add_run(seg.replace("<b>", "").replace("</b>", "")),
                 size=10.5, bold=bold)

    if caption:
        c = doc.add_paragraph()
        c.paragraph_format.space_before = Pt(4)
        c.paragraph_format.space_after = Pt(8)
        font(c.add_run(caption), size=11, italic=True)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def _clean_algo_cell(text: str) -> str:
    text = text.replace("\\quad", " ")
    return detex(text)


def _algo_code(text: str) -> str:
    """Clean one pseudocode line, keeping \\textbf{} keywords as <b> markers."""
    text = text.replace("\\quad", " ")
    text = re.sub(r"\\hspace\*\{[\d.]+em\}", "", text)
    text = re.sub(r"\\textbf\{([^{}]*)\}", r"<b>\1</b>", text)
    return detex(text)


def parse_algobox(block: str):
    """Parse one `% ALGOBOX-START ... % ALGOBOX-END` block (see algo_box() in
    build_paper_tex.py) into (caption, inputs, outputs, [steps]). The caption
    ("Algorithm N: Title") lives in its own \\begin{center} below the boxed
    tabular, which holds only the Input/Output lines and numbered steps."""
    cap_m = re.search(r"\\textbf\{(Algorithm\s*\d+:)\}\s*([^\n]*)", block)
    caption = _clean_algo_cell(f"{cap_m.group(1)} {cap_m.group(2)}") if cap_m else ""
    tab = re.search(r"\\begin\{tabular\}.*?\n(.+?)\\hline\s*\\end\{tabular\}",
                    block, re.S)
    rows_raw = tab.group(1) if tab else block
    rows = [r.strip() for r in rows_raw.split("\\\\") if r.strip()]
    inputs = outputs = ""
    steps = []
    for r in rows:
        c = _clean_algo_cell(r)
        if not c:
            continue
        if c.lower().startswith("input:"):
            inputs = c[len("input:"):].strip()
        elif c.lower().startswith("output:"):
            outputs = c[len("output:"):].strip()
        else:
            level = 0
            mh = re.search(r"\\hspace\*\{([\d.]+)em\}", r)
            if mh:
                level = int(round(float(mh.group(1)) / 1.5))
            code = _algo_code(r)
            m = re.match(r"^\d+\.\s*(.*)$", code)
            steps.append((level, m.group(1) if m else code))
    return caption, inputs, outputs, steps


# ----------------------------------------------------------------------------- #
#  main conversion
# ----------------------------------------------------------------------------- #
def parse_tex(tex: str):
    """Yield ('kind', payload) blocks in document order."""
    body = tex.split(r"\thispagestyle{plain}", 1)[1]
    body = body.split(r"\end{document}")[0]

    # --- title / authors (fixed strings; not worth parsing the centred block) ---
    tm = re.search(r"bfseries (.+?)\\par", body, re.S)
    title = re.sub(r"\s+", " ",
                   re.sub(r"\\\\\[[^\]]*\]", " ", tm.group(1))).strip() if tm else \
        "SafeNodule-MAS: An Instrumented Multi-Agent Architecture for Safe, " \
        "Calibrated, and Cost-Efficient Lung Nodule Diagnosis"
    yield ("title", (
        title,
        "Lochan Narayan K\u00b9      Manikandan R\u00b9,*",
        "\u00b9School of Computing, SASTRA Deemed University, Thanjavur, "
        "Tamil Nadu, India\n"
        "lochannarayan11@gmail.com      manikandan75@core.sastra.edu      "
        "(* corresponding author)"))

    # --- abstract ---
    ab = re.search(r"\\begin\{adjustwidth\}.*?selectfont\\justifying (.+?)\\par\}",
                   body, re.S)
    if ab:
        yield ("abstract", detex(ab.group(1)))
    kw = re.search(r"\\textbf\{Keywords:\}\s*(.+?)\\par", body, re.S)
    if kw:
        yield ("keywords", detex(kw.group(1)))

    # --- walk the rest linearly ---
    rest = body[body.find(r"\section{"):]
    i = 0
    fig_n = 0
    tab_n = 0
    tok = re.compile(
        r"\\section\{(?P<sec>[^}]*)\}"
        r"|\\section\*\{(?P<secstar>[^}]*)\}"
        r"|\\subsection\{(?P<sub>[^}]*)\}"
        r"|\\subsection\*\{(?P<substar>[^}]*)\}"
        r"|%\s*ALGOBOX-START(?P<algobox>.*?)%\s*ALGOBOX-END"
        r"|\\begin\{(?P<envs>table|figure\*?)\}(?P<env>.*?)\\end\{(?P=envs)\}"
        r"|\\begin\{thebibliography\}(?P<bib>.*?)\\end\{thebibliography\}", re.S)
    for m in tok.finditer(rest):
        pre = rest[i:m.start()]
        i = m.end()
        for chunk in re.split(r"\n\s*\n", pre):
            chunk = chunk.strip()
            if not chunk or chunk.startswith("%") or chunk.startswith("\\"):
                if not (chunk and re.match(r"\\(emph|textbf|textit|cite)", chunk)):
                    continue
            txt = detex(chunk)
            if txt:
                yield ("para", txt)
        if m.group("sec") is not None:
            yield ("h1", detex(m.group("sec")))
        elif m.group("secstar") is not None:
            yield ("h1star", detex(m.group("secstar")))
        elif m.group("sub") is not None:
            yield ("h2", detex(m.group("sub")))
        elif m.group("substar") is not None:
            yield ("h2star", detex(m.group("substar")))
        elif m.group("algobox") is not None:
            yield ("algo", parse_algobox(m.group("algobox")))
        elif m.group("envs") and m.group("envs").startswith("figure"):
            env = m.group("env")
            fig_n += 1
            fn = re.search(r"includegraphics\[[^\]]*\]\{figures/([^}]+)\}", env)
            cap = re.search(r"\\caption\{(.+?)\}\s*\n?\s*\\label", env, re.S)
            wide = m.group("envs") == "figure*"
            cap_txt = detex(cap.group(1)) if cap else ""
            yield ("fig", (fn.group(1) if fn else None,
                           f"Figure {fig_n}: {cap_txt}", wide))
        elif m.group("envs") == "table":
            env = m.group("env")
            tab_n += 1
            cap = re.search(r"\\caption\{(.+?)\}\s*\n\s*\\label", env, re.S)
            cap_txt = detex(cap.group(1)) if cap else ""
            cap = f"Table {tab_n}: {cap_txt}"
            tab = re.search(
                r"\\begin\{tabular\}\{(?:[^{}]|\{[^{}]*\})*\}(.+?)\\end\{tabular\}",
                env, re.S)
            headers, rows = [], []
            if tab:
                raw_rows = [r.strip() for r in tab.group(1).split(r"\\")]
                data = []
                for r in raw_rows:
                    r = re.sub(r"\\hline|\\rowcolor\{[^}]*\}", "", r).strip()
                    if not r:
                        continue
                    data.append([detex(c) for c in r.split("&")])
                if data:
                    headers, rows = data[0], data[1:]
            yield ("table", (cap, headers, rows))
        elif m.group("bib") is not None:
            entries = re.findall(r"\\bibitem\{[^}]*\}\s*(.+?)(?=\\bibitem|\Z)",
                                 m.group("bib"), re.S)
            yield ("bib", [_bib_clean(e) for e in entries])

    tail = rest[i:]
    for chunk in re.split(r"\n\s*\n", tail):
        chunk = chunk.strip()
        if chunk and not chunk.startswith("\\"):
            t = detex(chunk)
            if t:
                yield ("para", t)


def build():
    tex = TEX.read_text(encoding="utf-8")
    doc = Document()
    sec = doc.sections[0]
    for s in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, s, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    # page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(); font(r, size=BODY_SIZE)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

    prev_heading = False
    h1_num = 0
    h2_num = 0
    for kind, payload in parse_tex(tex):
        if kind == "title":
            title, authors, affil = payload
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            font(p.add_run(title), size=20, bold=True)
            if authors:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(6)
                font(p.add_run(authors), size=13)
            if affil:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(14)
                for li, line_ in enumerate(affil.split("\n")):
                    r = p.add_run(line_)
                    font(r, size=11, italic=True)
                    if li == 0:
                        r.add_break()
        elif kind == "abstract":
            para(doc, "Abstract", size=13, bold=True, first_indent=False,
                 align=WD_ALIGN_PARAGRAPH.LEFT, after=8)
            para(doc, payload, size=BODY_SIZE, first_indent=False,
                 l_indent=0.3, r_indent=0.3, after=8)
            prev_heading = True
        elif kind == "keywords":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(18)
            font(p.add_run("Keywords: "), size=BODY_SIZE, bold=True)
            font(p.add_run(payload), size=BODY_SIZE)
        elif kind == "h1":
            h1_num += 1; h2_num = 0
            heading(doc, f"{h1_num}  {payload}", 1); prev_heading = True
        elif kind == "h1star":
            heading(doc, payload, 1); prev_heading = True
        elif kind == "h2":
            h2_num += 1
            heading(doc, f"{h1_num}.{h2_num}  {payload}", 2); prev_heading = True
        elif kind == "h2star":
            heading(doc, payload, 3); prev_heading = True
        elif kind == "para":
            para(doc, payload, first_indent=not prev_heading); prev_heading = False
        elif kind == "table":
            add_table(doc, payload[0], payload[1], payload[2]); prev_heading = False
        elif kind == "fig":
            fn, cap, wide = payload
            add_figure(doc, FIGS / fn if fn else Path("nope"),
                       cap, width=6.5 if wide else 6.0)
            prev_heading = False
        elif kind == "algo":
            algo_cap, inputs, outputs, steps = payload
            add_algorithm(doc, algo_cap, inputs, outputs, steps)
            prev_heading = False
        elif kind == "bib":
            heading(doc, "References", 1)
            for idx, e in enumerate(payload, 1):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Inches(0.35)
                pf.first_line_indent = Inches(-0.35)
                pf.space_after = Pt(4)
                pf.line_spacing = 1.0
                font(p.add_run(f"{idx}.  "), size=BODY_SIZE)
                for seg in re.split(r"(<i>.*?</i>)", e):
                    if not seg:
                        continue
                    it = seg.startswith("<i>")
                    font(p.add_run(seg.replace("<i>", "").replace("</i>", "")),
                         size=BODY_SIZE, italic=it)
            prev_heading = False

    OUT.parent.mkdir(exist_ok=True)
    doc.core_properties.title = ("SafeNodule-MAS: An Instrumented Multi-Agent "
                                 "Architecture for Safe, Calibrated, and "
                                 "Cost-Efficient Lung Nodule Diagnosis")
    doc.core_properties.author = "Lochan Narayan K; Manikandan R"
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
